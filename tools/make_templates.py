from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent.parent / "templates_docx"


def _setup(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)


def _p(doc: Document, text: str = "", bold: bool = False, align=None, after: int = 6, indent: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    return p


def _subsection(doc: Document, text: str, indent: bool = True) -> None:
    _p(doc, text, indent=indent)


def _parties_table(doc: Document, role_tag: str, label: str) -> None:
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [
        ("ФИО", f"{role_tag}_fio"),
        ("Паспорт", f"{role_tag}_passport_line"),
        ("Адрес", f"{role_tag}_address"),
    ]
    for i, (title, key) in enumerate(rows):
        table.rows[i].cells[0].text = title
        table.rows[i].cells[1].text = f"{{{{ {key} }}}}"
    _p(doc, "", after=2)


def _signatures(doc: Document) -> None:
    _p(doc, "", after=12)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for i, text in enumerate(["Продавец: _______________ / _______________ /", "Покупатель: _______________ / _______________ /"]):
        table.rows[i].cells[0].text = text
        table.rows[i].cells[1].text = text


def _parties_block(doc: Document, seller_label: str, buyer_label: str) -> None:
    text = (
        f"«{seller_label}» — Продавец: {{{{ seller_fio }}}}, {{{{ seller_birth }}}} года рождения, "
        f"паспорт {{{{ seller_passport_line }}}}, зарегистрированный(ая) по адресу: {{{{ seller_address }}}}, "
        f"тел.: {{{{ seller_phone }}}}, именуемый(ая) в дальнейшем «Продавец», и "
        f"«{buyer_label}» — Покупатель: {{{{ buyer_fio }}}}, {{{{ buyer_birth }}}} года рождения, "
        f"паспорт {{{{ buyer_passport_line }}}}, зарегистрированный(ая) по адресу: {{{{ buyer_address }}}}, "
        f"тел.: {{{{ buyer_phone }}}}, именуемый(ая) в дальнейшем «Покупатель», "
        f"совместно именуемые «Стороны», заключили настоящий Договор о нижеследующем:"
    )
    _p(doc, text, indent=True)


def build_movables() -> None:
    doc = Document()
    _setup(doc)
    _p(doc, "ДОГОВОР КУПЛИ-ПРОДАЖИ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    _p(doc, "движимого имущества", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    _p(doc, "г. {{ city }}", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    _p(doc, "«____» ________________ {{ signing_date }} года", align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)

    _parties_block(doc, "Продавец", "Покупатель")

    _p(doc, "1. Предмет договора", bold=True)
    _subsection(doc, "1.1. Продавец обязуется передать в собственность Покупателя, а Покупатель обязуется принять и оплатить следующее имущество (далее — «Товар»): {{ item_name }}.")
    _subsection(doc, "1.2. Характеристики и состояние Товара: {{ item_desc }}.")
    _subsection(doc, "1.3. Товар принадлежит Продавцу на праве собственности, не заложен, не арестован, в споре и под запрещением не состоит.")

    _p(doc, "2. Цена и порядок расчётов", bold=True)
    _subsection(doc, "2.1. Цена Товара составляет {{ price_digits }} рублей ({{ price_words }}). НДС не облагается.")
    _subsection(doc, "2.2. Оплата производится {{ payment_method }} в срок: {{ payment_term }}.")

    _p(doc, "3. Передача товара", bold=True)
    _subsection(doc, "3.1. Продавец обязуется передать Товар Покупателю в срок: {{ transfer_term }}.")
    _subsection(doc, "3.2. Передача Товара осуществляется по факту его вручения Покупателю. Риск случайной гибели или повреждения Товара переходит к Покупателю с момента передачи.")

    _p(doc, "4. Ответственность сторон", bold=True)
    _subsection(doc, "4.1. За неисполнение или ненадлежащее исполнение обязательств по настоящему Договору Стороны несут ответственность в соответствии с законодательством Российской Федерации.")
    _subsection(doc, "4.2. Споры и разногласия Стороны разрешают путём переговоров, а при недостижении согласия — в судебном порядке в соответствии с законодательством Российской Федерации.")

    _p(doc, "5. Прочие условия", bold=True)
    _subsection(doc, "5.1. Настоящий Договор вступает в силу с момента его подписания и действует до полного исполнения Сторонами своих обязательств.")
    _subsection(doc, "5.2. Договор составлен в двух экземплярах, имеющих равную юридическую силу, по одному для каждой из Сторон.")

    _p(doc, "6. Данные Сторон", bold=True)
    _parties_table(doc, "seller", "Продавец")
    _parties_table(doc, "buyer", "Покупатель")

    _p(doc, "7. Подписи Сторон", bold=True)
    _signatures(doc)

    doc.save(OUT / "movables.docx")
    print("movables.docx OK")


def build_auto() -> None:
    doc = Document()
    _setup(doc)
    _p(doc, "ДОГОВОР КУПЛИ-ПРОДАЖИ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    _p(doc, "транспортного средства", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    _p(doc, "г. {{ city }}", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    _p(doc, "«____» ________________ {{ signing_date }} года", align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)

    _parties_block(doc, "Продавец", "Покупатель")

    _p(doc, "1. Предмет договора", bold=True)
    _subsection(doc, "1.1. Продавец обязуется передать в собственность Покупателя, а Покупатель обязуется принять и оплатить транспортное средство (далее — «ТС»):")
    _subsection(doc, "1.2. Характеристики ТС: марка, модель — {{ car_make_model }}; год выпуска — {{ car_year }}; VIN — {{ car_vin }}; цвет кузова — {{ car_body_color }}; государственный регистрационный знак — {{ car_regno }}; пробег — {{ car_mileage }} км.")
    _subsection(doc, "1.3. Паспорт транспортного средства (ПТС): серия {{ car_pts_series }}, номер {{ car_pts_number }}, выдан {{ car_pts_date }}.")
    _subsection(doc, "1.4. Продавец гарантирует, что ТС не находится в залоге, под арестом, не является предметом спора и не имеет иных обременений.")

    _p(doc, "2. Цена и порядок расчётов", bold=True)
    _subsection(doc, "2.1. Цена ТС составляет {{ price_digits }} рублей ({{ price_words }}). НДС не облагается.")
    _subsection(doc, "2.2. Оплата производится {{ payment_method }} в срок: {{ payment_term }}.")

    _p(doc, "3. Передача ТС", bold=True)
    _subsection(doc, "3.1. Продавец обязуется передать ТС Покупателю вместе с ПТС в срок: {{ transfer_term }}.")
    _subsection(doc, "3.2. Передача ТС осуществляется по факту вручения ключей и документов. Риск случайной гибели или повреждения ТС переходит к Покупателю с момента передачи.")
    _subsection(doc, "3.3. Снятие и постановка ТС на регистрационный учёт осуществляются Покупателем в соответствии с Правилами регистрации транспортных средств.")

    _p(doc, "4. Ответственность сторон", bold=True)
    _subsection(doc, "4.1. За неисполнение или ненадлежащее исполнение обязательств по настоящему Договору Стороны несут ответственность в соответствии с законодательством Российской Федерации.")
    _subsection(doc, "4.2. Споры и разногласия Стороны разрешают путём переговоров, а при недостижении согласия — в судебном порядке в соответствии с законодательством Российской Федерации.")

    _p(doc, "5. Прочие условия", bold=True)
    _subsection(doc, "5.1. Настоящий Договор вступает в силу с момента его подписания и действует до полного исполнения Сторонами своих обязательств.")
    _subsection(doc, "5.2. Договор составлен в трёх экземплярах, имеющих равную юридическую силу: по одному для каждой из Сторон и один — для регистрационного подразделения ГИБДД.")

    _p(doc, "6. Данные Сторон", bold=True)
    _parties_table(doc, "seller", "Продавец")
    _parties_table(doc, "buyer", "Покупатель")

    _p(doc, "7. Подписи Сторон", bold=True)
    _signatures(doc)

    doc.save(OUT / "auto.docx")
    print("auto.docx OK")


def build_realty() -> None:
    doc = Document()
    _setup(doc)
    _p(doc, "ДОГОВОР КУПЛИ-ПРОДАЖИ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    _p(doc, "недвижимого имущества", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    _p(doc, "г. {{ city }}", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    _p(doc, "«____» ________________ {{ signing_date }} года", align=WD_ALIGN_PARAGRAPH.RIGHT, after=10)

    _parties_block(doc, "Продавец", "Покупатель")

    _p(doc, "1. Предмет договора", bold=True)
    _subsection(doc, "1.1. Продавец обязуется передать в собственность Покупателя, а Покупатель обязуется принять и оплатить следующий объект недвижимости (далее — «Объект»):")
    _subsection(doc, "1.2. Характеристики Объекта: адрес — {{ realty_address }}; кадастровый номер — {{ realty_cadastral }}; общая площадь — {{ realty_area }} кв. м; количество комнат — {{ realty_rooms }}; этаж — {{ realty_floor }} из {{ realty_floors }}; год постройки — {{ realty_year }}.")
    _subsection(doc, "1.3. Объект принадлежит Продавцу на праве собственности, что подтверждается: {{ realty_doc_num }}.")
    _subsection(doc, "1.4. Объект не заложен, не арестован, в споре и под запрещением не состоит, правами третьих лиц не обременён.")

    _p(doc, "2. Цена и порядок расчётов", bold=True)
    _subsection(doc, "2.1. Цена Объекта составляет {{ price_digits }} рублей ({{ price_words }}). НДС не облагается.")
    _subsection(doc, "2.2. Оплата производится {{ payment_method }} в срок: {{ payment_term }}.")

    _p(doc, "3. Передача Объекта", bold=True)
    _subsection(doc, "3.1. Продавец обязуется передать Объект Покупателю по передаточному акту в порядке и сроки, согласованные Сторонами, но не позднее дня государственной регистрации перехода права собственности.")
    _subsection(doc, "3.2. Обязательство Продавца по передаче Объекта считается исполненным после подписания передаточного акта обеими Сторонами.")

    _p(doc, "4. Государственная регистрация", bold=True)
    _subsection(doc, "4.1. Переход права собственности на Объект подлежит государственной регистрации в Федеральной службе государственной регистрации, кадастра и картографии (Росреестр).")
    _subsection(doc, "4.2. Расходы по государственной регистрации перехода права собственности Стороны несут в соответствии с законодательством.")

    _p(doc, "5. Ответственность сторон", bold=True)
    _subsection(doc, "5.1. За неисполнение или ненадлежащее исполнение обязательств по настоящему Договору Стороны несут ответственность в соответствии с законодательством Российской Федерации.")
    _subsection(doc, "5.2. Споры и разногласия Стороны разрешают путём переговоров, а при недостижении согласия — в судебном порядке в соответствии с законодательством Российской Федерации.")

    _p(doc, "6. Прочие условия", bold=True)
    _subsection(doc, "6.1. Настоящий Договор вступает в силу с момента его подписания и действует до полного исполнения Сторонами своих обязательств.")
    _subsection(doc, "6.2. Договор составлен в трёх экземплярах, имеющих равную юридическую силу: по одному для каждой из Сторон и один — для регистрирующего органа.")

    _p(doc, "7. Данные Сторон", bold=True)
    _parties_table(doc, "seller", "Продавец")
    _parties_table(doc, "buyer", "Покупатель")

    _p(doc, "8. Подписи Сторон", bold=True)
    _signatures(doc)

    doc.save(OUT / "realty.docx")
    print("realty.docx OK")


OUT.mkdir(parents=True, exist_ok=True)

if __name__ == "__main__":
    build_movables()
    build_auto()
    build_realty()